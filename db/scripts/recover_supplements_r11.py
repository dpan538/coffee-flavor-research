#!/usr/bin/env python3
"""Recover structured R11 tables from sanctioned Europe PMC supplements.

The script reads only the frozen R10 candidate list and R11 first-pass report.
It parses CSV/TSV/XLSX/XML/HTML members from Europe PMC supplementary ZIPs.
PDF and legacy binary spreadsheet members are inventoried but never admitted.
No estimator is constructed and no parameter is fitted.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import time
import zipfile

from openpyxl import load_workbook

from extract_candidates_r11 import cache_fetch, digest, parse_html, parse_xml, stable_bytes
from stratify_extraction_r11 import (
    R10,
    R11,
    load_direct_registry,
    ontology_records,
    sample_axis_records,
)


DEFAULT_CACHE = Path("/private/tmp/coffee-flavor-r11-supplement-cache")
DEFAULT_RAW = R11 / "extraction_report.json"
DEFAULT_OUTPUT = R11 / "supplement_recovery.json"
MEMBER_LIMIT = 25_000_000


def tabular_members(name: str, body: bytes) -> list[dict]:
    suffix = Path(name).suffix.lower()
    tables: list[dict] = []
    if suffix in {".csv", ".tsv", ".txt"}:
        text = body.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = [[str(cell).strip() for cell in row] for row in csv.reader(StringIO(text), dialect)]
        rows = [row for row in rows if any(row)]
        if rows:
            tables.append({"table_id": name, "label": name, "caption": name, "rows": rows})
    elif suffix == ".xlsx":
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = []
            for values in sheet.iter_rows(values_only=True):
                row = ["" if value is None else str(value).strip() for value in values]
                while row and not row[-1]:
                    row.pop()
                if any(row):
                    rows.append(row)
            if rows:
                label = f"{name}#{sheet.title}"
                tables.append({"table_id": label, "label": label, "caption": label, "rows": rows})
    elif suffix in {".xml", ".nxml"}:
        parsed, _ = parse_xml(body)
        for table in parsed:
            table["table_id"] = f"{name}#{table['table_id']}"
        tables.extend(parsed)
    elif suffix in {".html", ".htm"}:
        parsed, _ = parse_html(body)
        for table in parsed:
            table["table_id"] = f"{name}#{table['table_id']}"
        tables.extend(parsed)
    return tables


def primary_article_text(outcome: dict, primary_cache: Path) -> str:
    retrieval = outcome.get("retrieval") or {}
    requested = retrieval.get("requested_url")
    if not requested:
        return outcome.get("title", "")
    path = primary_cache / f"{digest(requested.encode())}.body"
    if not path.exists():
        return outcome.get("title", "")
    body = path.read_bytes()
    try:
        if body.lstrip().startswith(b"<?xml") or retrieval.get("content_type") in {
            "text/xml",
            "application/xml",
        }:
            _, metadata = parse_xml(body)
        else:
            _, metadata = parse_html(body)
        return metadata["plain_text"]
    except Exception:
        return outcome.get("title", "")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-report", type=Path, default=DEFAULT_RAW)
    parser.add_argument("--primary-cache", type=Path, default=Path("/private/tmp/coffee-flavor-r11-fulltext-cache"))
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--timeout", type=float, default=25.0)
    args = parser.parse_args()

    raw = json.loads(args.raw_report.read_text())
    candidates = {
        row["candidate_id"]: row
        for row in json.loads((R10 / "acquisition_manifest.json").read_text())["discovery_candidates"]
    }
    direct, rules = load_direct_registry()
    records = []
    outcomes = []
    for outcome in raw["candidate_outcomes"]:
        retrieval = outcome.get("retrieval") or {}
        match = re.search(r"/(PMC\d+)/fullTextXML", retrieval.get("final_url", ""), re.I)
        if not match:
            continue
        pmcid = match.group(1).upper()
        endpoint = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/supplementaryFiles"
        base = {
            "candidate_id": outcome["candidate_id"],
            "doi": outcome["doi"],
            "pmcid": pmcid,
            "endpoint": endpoint,
        }
        try:
            response = cache_fetch(endpoint, args.cache_dir, args.timeout)
            body = response.pop("body")
        except Exception as exc:
            outcomes.append(base | {"status": "SUPPLEMENT_RETRIEVAL_FAILED", "detail": f"{type(exc).__name__}:{str(exc)[:180]}"})
            continue
        if not zipfile.is_zipfile(BytesIO(body)):
            outcomes.append(base | {"status": "NO_OA_SUPPLEMENT_ZIP", "retrieval": response})
            continue
        candidate = candidates[outcome["candidate_id"]]
        article_text = primary_article_text(outcome, args.primary_cache)
        member_rows = []
        candidate_records = []
        try:
            with zipfile.ZipFile(BytesIO(body)) as archive:
                for member in archive.infolist():
                    suffix = Path(member.filename).suffix.lower()
                    item = {"name": member.filename, "bytes": member.file_size, "suffix": suffix}
                    if member.is_dir():
                        continue
                    if member.file_size > MEMBER_LIMIT:
                        member_rows.append(item | {"status": "MEMBER_TOO_LARGE_NOT_PARSED"})
                        continue
                    if suffix in {".pdf", ".xls", ".doc", ".docx"}:
                        member_rows.append(item | {"status": "UNSTRUCTURED_OR_LEGACY_NOT_ADMITTED"})
                        continue
                    try:
                        member_body = archive.read(member)
                        tables = tabular_members(member.filename, member_body)
                    except Exception as exc:
                        member_rows.append(item | {"status": "PARSE_FAILED", "detail": f"{type(exc).__name__}:{str(exc)[:180]}"})
                        continue
                    if not tables:
                        member_rows.append(item | {"status": "NO_SUPPORTED_STRUCTURED_TABLE"})
                        continue
                    supplement_outcome = dict(outcome)
                    supplement_outcome["retrieval"] = {
                        "requested_url": endpoint,
                        "final_url": endpoint + "#" + member.filename,
                        "sha256": digest(member_body),
                        "content_type": suffix.removeprefix("."),
                        "bytes": len(member_body),
                    }
                    produced = []
                    for table in tables:
                        produced.extend(ontology_records(candidate, supplement_outcome, table, direct, rules))
                        produced.extend(sample_axis_records(candidate, supplement_outcome, table, article_text, direct, rules))
                    produced = list({row["record_id"]: row for row in produced}.values())
                    candidate_records.extend(produced)
                    member_rows.append(item | {"status": "PARSED", "table_count": len(tables), "typed_record_count": len(produced), "sha256": digest(member_body)})
        except zipfile.BadZipFile as exc:
            outcomes.append(base | {"status": "SUPPLEMENT_ZIP_INVALID", "detail": str(exc)[:180], "retrieval": response})
            continue
        records.extend(candidate_records)
        outcomes.append(
            base
            | {
                "status": "RECOVERED_TYPED_RECORDS" if candidate_records else "NO_GOVERNED_RECORD_RECOVERED",
                "retrieval": response,
                "member_count": len(member_rows),
                "typed_record_count": len(candidate_records),
                "members": member_rows,
            }
        )

    records = sorted({row["record_id"]: row for row in records}.values(), key=lambda row: row["record_id"])
    report = {
        "contract_version": "r11.sanctioned-supplement-recovery.v1",
        "generated_at_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "retrieval_policy": [
            "EUROPE_PMC_SANCTIONED_SUPPLEMENT_ZIP",
            "CSV_TSV_XLSX_XML_HTML_ONLY",
            "PDF_AND_LEGACY_BINARY_NOT_ADMITTED",
        ],
        "guards": {
            "fit_count": 0,
            "real_label_metric_count": 0,
            "new_mapping_rules_created": 0,
            "pdf_records_admitted": 0,
        },
        "summary": {
            "candidate_supplement_endpoints_attempted": len(outcomes),
            "outcome_counts": dict(sorted(Counter(row["status"] for row in outcomes).items())),
            "typed_record_count": len(records),
            "shape_record_counts": dict(sorted(Counter(row["shape_id"] for row in records).items())),
            "coffee_group_count": len({row.get("coffee_group_id") for row in records if row.get("coffee_group_id")}),
        },
        "candidate_outcomes": outcomes,
        "records": records,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(stable_bytes(report))
    print(json.dumps(report["summary"], indent=2))


if __name__ == "__main__":
    main()
