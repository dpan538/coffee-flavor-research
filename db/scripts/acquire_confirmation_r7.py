"""Bounded R7 source intake and typed raw-file audit, without model evaluation.

Original files and parsed observations remain in owner storage. Metadata-only,
aggregate and one-material sources never become fresh confirmation examples.
The CLI distinguishes actual parsing from hash verification; neither fits models.
"""

from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import shutil
import subprocess
import tarfile
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from acquire_supervision_r5 import digest, fetch, now, save

VERSION = "m2-r7-confirmation-source-audit.v1"
R7 = Path("revisions/r7")
FIELDS = ("Salty", "Spice.Cabinet", "Sweet", "Bitter", "Nutty")
EXPOSITION_SHA = "f9c9f49ab94b0b5700e9db5a456ef279544c75e1d92ba55cebd4f8ff17de0bd3"
YUNNAN_SHA = "536b2cf59f038daa05f198bf336c61cbb816f288a207f5bdbe763fb6e826549f"
YUNNAN_XML_SHA = "c110c0ddf1f2b83269c11d5ba499f45d37153d86d90a4405c3f2981435ebf9a4"


def protocol():
    return {
        "version": VERSION,
        "purpose": "NEW_CONFIRMATION_SOURCE_AVAILABILITY_AUDIT",
        "model_evaluation_allowed": False,
        "training_allowed": False,
        "source_acquisition_block": "bounded targeted retrieval, approximately 2-3 hour cap; actual timestamps retained",
        "target": "new multi-coffee source or 15-30 traceable groups is a search goal, not an admission quota",
        "confirmation_gate": [
            "actual original observation rows, with rights applicable to this noncommercial research",
            "verified coffee/material/condition and duplicate dependency identity",
            "not previously used by base models, coverage or prior development",
            "disjoint original A, B and T observations; concept overlap remains legal",
            "source measurement and fixed MAP_BASE target compatibility reviewed",
            "owner freezes cohort/roles/T/metrics before any C00/C01 evaluation",
        ],
        "forbidden": [
            "renaming old anonymous samples or rotating graders to create new coffee groups",
            "reconstructing individuals from averages, frequencies, PCA or plots",
            "converting native numeric zero into universal sensory absence",
            "using confirmation targets to select sources or modify mappings",
            "author requests, recruitment, fees or access-control bypass",
            "unrelated score-recovery models",
        ],
        "typed_only": ["EXPOSITION_SINGLE_MATERIAL", "YUNNAN_AGGREGATE_M"],
        "public_policy": "aggregate source facts and owner-relative evidence hashes only; no answers or participant/sample identifiers",
    }


def checked_bytes(path, expected=None):
    raw = Path(path).read_bytes()
    if expected is not None and digest(raw) != expected:
        raise ValueError("SOURCE_HASH_CHANGED:" + Path(path).name)
    return raw


def parse_exposition_rows(rows):
    """Native long-form rows from the R array; no assumption about code 0."""
    observations = {}
    for row in rows:
        condition, participant, attribute = (
            row["condition"],
            row["participant"],
            row["attribute"],
        )
        if attribute not in FIELDS:
            raise ValueError("UNREGISTERED_EXPOSITION_ATTRIBUTE")
        value = float(row["value"])
        if not math.isfinite(value) or value not in (0, 1, 2):
            raise ValueError("EXPOSITION_NATIVE_CODE_REQUIRED")
        key = (condition, participant)
        observation = observations.setdefault(key, {})
        if attribute in observation:
            raise ValueError("DUPLICATE_SOURCE_OBSERVATION_CELL")
        observation[attribute] = int(value)
    if any(set(values) != set(FIELDS) for values in observations.values()):
        raise ValueError("INCOMPLETE_SOURCE_OBSERVATION")
    records = [
        {
            "material_id": "EXPOSITION_ONE_HONDURAN_COFFEE",
            "condition_id": key[0],
            "participant_id": key[1],
            "observation_id": "exposition:" + key[0] + ":" + key[1],
            "native_recorded_ratings": values,
            "observation_masks": {field: True for field in FIELDS},
        }
        for key, values in sorted(observations.items())
    ]
    return records


def parse_exposition(path):
    """Read the actual GPL-2 package data, excluding preference responses."""
    raw = checked_bytes(path, EXPOSITION_SHA)
    with tarfile.open(fileobj=io.BytesIO(raw), mode="r:gz") as archive:
        member = archive.getmember("ExPosition/data/coffee.data.rda")
        if member.size > 1_000_000:
            raise ValueError("UNEXPECTED_RDA_SIZE")
        rda = archive.extractfile(member).read()
        description = archive.extractfile("ExPosition/DESCRIPTION").read().decode()
        documentation = (
            archive.extractfile("ExPosition/man/coffee.data.Rd").read().decode()
        )
    if "License: GPL-2" not in description or "Honduran" not in documentation:
        raise ValueError("EXPOSITION_PROVENANCE_CHANGED")
    rscript = shutil.which("Rscript")
    if rscript is None:
        raise RuntimeError("RSCRIPT_REQUIRED_FOR_ACTUAL_RDA_PARSE")
    with tempfile.TemporaryDirectory(prefix="coffee-r7-rda-") as temp:
        temp = Path(temp)
        (temp / "source.rda").write_bytes(rda)
        (temp / "read.R").write_text(
            "args <- commandArgs(trailingOnly=TRUE)\n"
            "e <- new.env(); load(args[1], envir=e)\n"
            "x <- e$coffee.data$ratings\n"
            "stopifnot(identical(dim(x),c(4L,5L,10L)), !anyNA(x))\n"
            'd <- as.data.frame.table(x, responseName="value")\n'
            'names(d) <- c("condition","attribute","participant","value")\n'
            'write.table(d,args[2],sep="\\t",row.names=FALSE,quote=TRUE)\n'
        )
        subprocess.run(
            [
                rscript,
                str(temp / "read.R"),
                str(temp / "source.rda"),
                str(temp / "rows.tsv"),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        records = parse_exposition_rows(
            csv.DictReader(io.StringIO((temp / "rows.tsv").read_text()), delimiter="\t")
        )
    counts = Counter(
        v for row in records for v in row["native_recorded_ratings"].values()
    )
    if len(records) != 40 or counts != Counter({0: 73, 1: 95, 2: 32}):
        raise ValueError("EXPOSITION_ACTUAL_DATA_COUNT_CHANGED")
    return {
        "source_id": "EXPOSITION_SINGLE_MATERIAL",
        "source_sha256": digest(raw),
        "rda_sha256": digest(rda),
        "measurement_type": "SOURCE_NATIVE_0_2_RECORDED_ATTRIBUTE_RATING",
        "zero_semantics": "Native code retained; detailed anchors not verified; not CATA absence or a fine positive-mention label",
        "source_units": {
            "coffee_materials": 1,
            "conditions": 4,
            "participants": 10,
            "original_observations": 40,
            "observed_cells": 200,
        },
        "fields": list(FIELDS),
        "native_code_counts": dict(sorted(counts.items())),
        "excluded_fields": [
            "preferences (hedonic Yes/Impartial response, not sensory target)"
        ],
        "status": "TYPED_MEASUREMENT_ONLY_SINGLE_MATERIAL_NOT_MAIN_CONFIRMATION",
        "role_assignment": None,
        "confirmation_groups": 0,
        "records": records,
    }


def parse_yunnan(path, article_path):
    """The supplement is 25 product aggregates, not 10 original judges."""
    import openpyxl

    raw = checked_bytes(path, YUNNAN_SHA)
    article = checked_bytes(article_path, YUNNAN_XML_SHA)
    root = ET.fromstring(article)
    table = next(x for x in root.iter("table-wrap") if x.get("id") == "tbl1")
    metadata = {}
    for tr in table.iter("tr"):
        cells = [
            " ".join("".join(x.itertext()).split())
            for x in list(tr)
            if x.tag in ("td", "th")
        ]
        if len(cells) == 5 and cells[0] != "ID":
            metadata[cells[0]] = dict(
                zip(
                    (
                        "sample_id",
                        "growing_area",
                        "source_name",
                        "source_roast",
                        "manufacturer",
                    ),
                    cells,
                    strict=True,
                )
            )
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        workbook_raw = archive.read("mmc1.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_raw), data_only=True)
    sheet = workbook["Mi_individual_coffee_filtered"]
    fields = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))[1:]
    if len(fields) != 57 or len(set(fields)) != 57:
        raise ValueError("YUNNAN_REGISTERED_FIELD_COUNT_CHANGED")
    records = []
    for source_row, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if str(row[0]).startswith("*Please see"):
            if any(x is not None for x in row[1:]):
                raise ValueError("UNEXPECTED_FOOTNOTE_VALUES")
            continue
        if row[0] not in metadata:
            raise ValueError("YUNNAN_PRODUCT_ID_JOIN_FAILED")
        values = {}
        for field, value in zip(fields, row[1:], strict=True):
            if (
                not isinstance(value, (int, float))
                or not math.isfinite(value)
                or not 0 <= value <= 1
            ):
                raise ValueError("YUNNAN_AGGREGATE_M_VALUE_REQUIRED")
            values[field] = float(value)
        records.append(
            {**metadata[row[0]], "source_row": source_row, "aggregate_M_values": values}
        )
    if (
        len(records) != 25
        or len({r["sample_id"] for r in records}) != 25
        or set(metadata) != {r["sample_id"] for r in records}
    ):
        raise ValueError("YUNNAN_ACTUAL_PRODUCT_COUNT_CHANGED")
    return {
        "source_id": "YUNNAN_AGGREGATE_M",
        "source_sha256": digest(raw),
        "workbook_sha256": digest(workbook_raw),
        "article_sha256": digest(article),
        "measurement_type": "AGGREGATE_SQRT_FREQUENCY_TIMES_RELATIVE_INTENSITY_M",
        "source_units": {
            "named_products": 25,
            "growing_regions": 4,
            "published_aggregate_rows": 25,
            "aggregate_cells": 1425,
            "original_grader_rows_obtained": 0,
        },
        "field_count": 57,
        "selection_limit": "Source filtered 66 descriptors to 57 using aggregate M>5%; this is not a complete raw ballot",
        "individual_reconstruction_allowed": False,
        "zero_semantics": "Aggregate M zero, not a negative observation from an individual",
        "rights": "CC-BY-NC-ND-4.0; original and private inspection retained, no public derivative matrix",
        "status": "TYPED_AGGREGATE_ONLY_NO_DISJOINT_A_B_T",
        "role_assignment": None,
        "confirmation_groups": 0,
        "records": records,
    }


def zenodo_version_audit(source_dir, original_workbook):
    source_dir = Path(source_dir)
    current = json.loads((source_dir / "zenodo-current.json").read_text())
    latest = json.loads((source_dir / "zenodo-latest.json").read_text())
    versions = json.loads((source_dir / "zenodo-versions.json").read_text())
    file = next(x for x in current["files"] if x["key"] == "panelists_scores_EN.xlsx")
    original = Path(original_workbook).read_bytes()
    equal = file["checksum"] == "md5:" + hashlib.md5(original).hexdigest()
    unchanged = str(current["id"]) == str(latest["id"]) == "20840464" and equal
    if not unchanged:
        raise ValueError("SOURCE_VERSION_CHANGED_REQUIRES_FRESH_IDENTITY_REVIEW")
    return {
        "source_id": "ZENODO_VERSION_AND_AUTHOR_REPOSITORY",
        "latest_record_id": str(latest["id"]),
        "version_count": versions["hits"]["total"],
        "panel_workbook_unchanged": equal,
        "original_workbook_sha256": digest(original),
        "source_samples_previously_inspected": 196,
        "source_grader_rows_previously_inspected": 526,
        "new_original_rows": 0,
        "new_coffee_groups": 0,
        "status": "NO_NEW_VERSION_OR_NEW_OBSERVATIONS",
        "anonymous_and_historical_rows_not_relabelled_fresh": True,
    }


def public_typed_summary(payload):
    """An allowlist prevents raw source identifiers or responses being published."""
    keys = (
        "source_id",
        "source_sha256",
        "measurement_type",
        "source_units",
        "field_count",
        "zero_semantics",
        "selection_limit",
        "rights",
        "status",
        "confirmation_groups",
    )
    return {key: payload[key] for key in keys if key in payload}


def source_routes():
    """Admission decisions derive from source identity/schema/rights, not scores."""
    return [
        {
            "source_id": "EXPOSITION_SINGLE_MATERIAL",
            "doi_or_url": "https://cran.r-project.org/src/contrib/ExPosition_2.11.0.tar.gz",
            "files": [
                "ExPosition_2.11.0.tar.gz",
                "ExPosition_DESCRIPTION.txt",
                "ExPosition_coffee.data.Rd",
            ],
            "rights": "GPL-2 package",
            "status": "TYPED_ONLY",
            "reason": "Actual 1 coffee, 4 conditions, 10 participants, 200 recorded attribute cells; not multicoffee or fixed fine-mention targets.",
        },
        {
            "source_id": "YUNNAN_AGGREGATE_M",
            "doi_or_url": "https://doi.org/10.1016/j.crfs.2022.07.010",
            "files": ["yunnan2022.xml", "yunnan2022-supp.zip"],
            "rights": "CC-BY-NC-ND-4.0",
            "status": "TYPED_ONLY",
            "reason": "Actual 25 named products by 57 source-filtered aggregate M columns, no original grader observations; no independent A/B/T.",
        },
        {
            "source_id": "NOVAK_REPOSITORY_CASE_EXAMPLES",
            "doi_or_url": "https://doi.org/10.5281/zenodo.17178120",
            "files": ["novak-sensory-economic.json", "novak-sensory-economic.pdf"],
            "rights": "CC-BY-4.0 repository",
            "status": "NO_VERIFIABLE_ORIGINAL_OBSERVATIONS",
            "reason": "Actual only 14-page framework PDF with illustrative A/B/C case tables; original coffee/panel/collection identities not documented. Dataset metadata label does not establish observations.",
        },
        {
            "source_id": "MATAS_MINAS_19928981",
            "doi_or_url": "https://api.figshare.com/v2/articles/19928981",
            "files": ["matas-minas-figshare.json"],
            "rights": "CC-BY-4.0 metadata",
            "status": "NO_ORIGINAL_PANEL_MATRIX",
            "reason": "Actual deposit inventory has two JPG files, no original coffee-by-grader table. No image digitization or invented people.",
        },
        {
            "source_id": "ELECTROCHEMICAL_APPRAISAL_30811541",
            "doi_or_url": "https://api.figshare.com/v2/articles/30811541",
            "files": ["electrochemical-appraisal-figshare.json"],
            "rights": "CC-BY-4.0",
            "status": "CHEMISTRY_ONLY_NOT_SENSORY_CONFIRMATION",
            "reason": "Actual file inventory describes electrochemical figures and chemistry workbooks, not independent human A/B/T. No chemical quality prediction branch added.",
        },
        {
            "source_id": "FRUIT_INFUSION_30946856",
            "doi_or_url": "https://api.figshare.com/v2/articles/30946856",
            "files": ["fruit-infusion-figshare.json"],
            "rights": "CC-BY-4.0 metadata",
            "status": "NO_FILES_IN_ACTUAL_DEPOSIT",
            "reason": "Actual thesis metadata contains no downloadable data files or original rows.",
        },
        {
            "source_id": "AUSTRALIA_CROP_TO_CUP",
            "doi_or_url": "https://doi.org/10.25918/thesis.535",
            "files": ["australia-crop-cup-datacite.json", "australia-crop-cup.html"],
            "rights": "Open access plus author copyright; dataset reuse not established",
            "status": "NO_ORIGINAL_MATRIX_OR_CONFIRMED_DATA_RIGHTS",
            "reason": "Metadata describes multiple varieties but actual landing response is an application shell; no original grading matrix obtained. Thesis sample counts are not acquired observations.",
        },
        {
            "source_id": "ONLINE_CONSUMER_2024",
            "doi_or_url": "https://doi.org/10.1016/j.foodres.2024.114349",
            "files": ["online-consumer2024-crossref.json"],
            "rights": "Article rights do not supply withheld response data",
            "status": "DATA_ON_REQUEST_NOT_ACQUIRED",
            "reason": "Published data-availability statement says on request; four coffee RATA products are study counts, not obtained participant observations. No request sent.",
        },
        {
            "source_id": "AMAZONIAN_ROBUSTA_2024",
            "doi_or_url": "https://doi.org/10.3390/beverages10030057",
            "files": ["amazonian2024.pdf"],
            "rights": "CC-BY-4.0 article",
            "status": "ARTICLE_AGGREGATES_SUPPLEMENT_NOT_OBTAINED",
            "reason": "Actual author-institution PDF: 2 clones, each natural/fermented, 127 consumer study participants; article CATA frequencies aggregated. Supplement described as expert summary and demographic table; actual public supplement request HTTP403, no individual matrix acquired.",
        },
        {
            "source_id": "GEISHA_PANAMA_2025",
            "doi_or_url": "https://doi.org/10.1002/fsn3.71278",
            "files": ["geisha2025.xml", "geisha2025-supp.zip"],
            "rights": "CC-BY-4.0 article; raw data explicitly restricted for privacy/ethics",
            "status": "RESTRICTED_ORIGINAL_DATA_NOT_ACQUIRED",
            "reason": "Article reports 24 coffee samples, 24 consumers, triplicate RATA. Actual public supplement contains sample metadata and aggregate tests, not individual responses. No participant count promoted to acquired rows; no restricted-data request.",
        },
        {
            "source_id": "HOME_USE_2023",
            "doi_or_url": "https://doi.org/10.1016/j.foodqual.2023.104905",
            "files": ["home-use2023.xml"],
            "rights": "Conditional Elsevier COVID research permission, not assumed CC",
            "status": "DATA_ON_REQUEST_NOT_ACQUIRED",
            "reason": "Actual XML says data on request. Four coffee products across three separate consumer-testing populations have aggregate article results only; no original A/B/T.",
        },
        {
            "source_id": "FREE_COMMENT_2025",
            "doi_or_url": "https://doi.org/10.1016/j.foodqual.2024.105377",
            "files": ["free-comment2025-crossref.json", "free-comment2025-author.html"],
            "rights": "Original review/participant-data rights not established",
            "status": "NO_OPEN_ORIGINAL_MATRIX_OBTAINED",
            "reason": "Author repository exposes bibliographic record and publisher link, not the 105-consumer RATA rows or matched original review identities. No scraped anonymous reviews become independent confirmation.",
        },
    ]


def inspect_additional_workbooks(source_dir):
    """Read actual source schemas without converting preference or analyte trials to coffee T."""
    import openpyxl

    source_dir = Path(source_dir)
    taste = checked_bytes(
        source_dir / "taste-sensitivity2018.xlsx",
        "1ef4bc85e81ff56e396c98108373a0b4626f89dc9c920f521b001979c590bab3",
    )
    rows = list(
        openpyxl.load_workbook(io.BytesIO(taste), data_only=True)["raw data"].values
    )
    if len(rows) != 94 or rows[0][:5] != (
        "ID",
        "Tasting note (information treatment)",
        "Bitterness",
        "Overall taste",
        "Preference level",
    ):
        raise ValueError("TASTE_SENSITIVITY_SCHEMA_CHANGED")
    data = [r for r in rows[1:] if r[0] is not None]
    if len(data) != 93 or len({r[0] for r in data}) != 90:
        raise ValueError("TASTE_SENSITIVITY_IDENTITY_CHANGED")
    sensory = checked_bytes(
        source_dir / "mozambioside-human-sensory.xlsx",
        "928f72702ff8c96f27d35f90cd807def47b6a8fa17ce0b9d19a44d8adbd09b38",
    )
    workbook = openpyxl.load_workbook(io.BytesIO(sensory), data_only=True)
    analytes = [s for s in workbook if s.title.startswith("analyte ")]
    if len(analytes) != 7 or any(
        s.cell(2, 2).value != "concentration [µM]" for s in analytes
    ):
        raise ValueError("ANALYTE_SENSORY_SCHEMA_CHANGED")
    return [
        {
            "source_id": "TASTE_SENSITIVITY_2018",
            "original_respondent_rows": len(data),
            "unique_recorded_ids": len({r[0] for r in data}),
            "identity_limitation": "93 rows but 90 recorded IDs; three IDs repeat. Do not infer 93 independent participants or coffee groups.",
            "columns": len(rows[0]),
            "treatment_counts": dict(Counter(r[1] for r in data)),
            "status": "REAL_RESPONSES_NO_TRACEABLE_MULTICOFFEE_A_B_T",
            "reason": "Actual bitterness plus taste/preference/purchase/PROP and demographic fields; no coffee identity column or independent fine-description reference. No hedonic target substitution.",
        },
        {
            "source_id": "MOZAMBIOSIDE_2024",
            "original_analyte_sheets": len(analytes),
            "status": "COMPOUND_CONCENTRATION_RECOGNITION_NOT_COFFEE_A_B_T",
            "reason": "Actual analyte concentration, genotype and native x/0/missing responses. These are compound trials, not seven coffee products; no universal sensory absence recoding or participant identity invention.",
        },
    ]


def file_evidence(source_dir, filename):
    path = Path(source_dir) / filename
    raw = checked_bytes(path)
    receipt = json.loads(path.with_name(path.name + ".acquisition.json").read_text())
    if receipt.get("sha256") != digest(raw) or receipt.get("status") != 200:
        raise ValueError("SOURCE_RECEIPT_MISMATCH:" + filename)
    return {
        "owner_relative_path": str(R7 / "sources" / filename),
        "sha256": digest(raw),
        "bytes": len(raw),
        "retrieved_utc": receipt["retrieved_utc"],
        "url": receipt["requested_url"],
    }


def build_report(owner):
    owner = Path(owner)
    source_dir = owner / R7 / "sources"
    original = owner.parent / "backend-model-20260905/sources/zenodo-panelists.xlsx"
    if not original.exists():
        raise FileNotFoundError("ORIGINAL_ZENODO_WORKBOOK_REQUIRED")
    parsed = [
        parse_exposition(source_dir / "ExPosition_2.11.0.tar.gz"),
        parse_yunnan(source_dir / "yunnan2022-supp.zip", source_dir / "yunnan2022.xml"),
    ]
    parsed_files = []
    for payload in parsed:
        filename = payload["source_id"].lower() + ".typed.private.json"
        receipt = save(owner / R7 / filename, payload)
        parsed_files.append({"owner_relative_path": str(R7 / filename), **receipt})
    routes = []
    for route in source_routes():
        evidence = [file_evidence(source_dir, f) for f in route["files"]]
        routes.append(
            {
                **{key: value for key, value in route.items() if key != "files"},
                "evidence": evidence,
                "admitted_confirmation_groups": 0,
                "model_evaluations": 0,
            }
        )
    additional = inspect_additional_workbooks(source_dir)
    for source_id, doi, files, reason in [
        (
            "TASTE_SENSITIVITY_2018",
            "10.17632/3yv7bdrczd.1",
            [
                "taste-sensitivity2018.html",
                "3yv7bdrczd-files.json",
                "taste-sensitivity2018.xlsx",
            ],
            additional[0]["reason"],
        ),
        (
            "MOZAMBIOSIDE_2024",
            "10.17632/xzjppbmn58.1",
            [
                "mozambioside2024.html",
                "xzjppbmn58-files.json",
                "mozambioside-human-sensory.xlsx",
            ],
            additional[1]["reason"],
        ),
    ]:
        routes.append(
            {
                "source_id": source_id,
                "doi_or_url": "https://doi.org/" + doi,
                "rights": "CC-BY-4.0 dataset",
                "status": "ORIGINAL_WORKBOOK_PARSED_NOT_MAIN_CONFIRMATION",
                "reason": reason,
                "evidence": [file_evidence(source_dir, f) for f in files],
                "admitted_confirmation_groups": 0,
                "model_evaluations": 0,
            }
        )
    versions = zenodo_version_audit(source_dir, original)
    version_evidence = [
        file_evidence(source_dir, f)
        for f in (
            "zenodo-current.json",
            "zenodo-latest.json",
            "zenodo-versions.json",
            "zenodo-author-records.json",
            "zenodo-author-repo-contents.json",
            "zenodo-author-repo-commits.json",
        )
    ]
    dedup = {
        "old_routes_not_counted_new": [
            "COTTER_B8993H_AND_2026_AI_REUSE",
            "CROIJMANS_DANS_ZKE_2WGQ",
            "INERA_FIGSHARE_25735122",
            "ROCCHETTI_PMC7736008",
            "R2_CROSSBRAND_APP15020948",
            "R2_REMOTE_PMC8548442",
        ],
        "accidental_rediscovery": "R2 remote-testing and cross-brand article identities were recognized against the old ledger after one new article fetch; no repeated supplement or restricted-data request followed.",
        "new_remote_zenodo_14712259_scope": "Actual new deposit is hemp-seed oil only, not the coffee arm; not imported as coffee evidence.",
        "remote_hemp_metadata": file_evidence(
            source_dir, "remote-testing-hemp-only.json"
        ),
        "other_old_failed_routes": "R5 ledger excludes NIR403, forbidden/no-sharing-rights source, old Dryad restricted route, R2 blank/proprietary forms; no access retry or identity rename.",
    }
    first = min(x["retrieved_utc"] for route in routes for x in route["evidence"])
    first = min(first, *(x["retrieved_utc"] for x in version_evidence))
    report = {
        "version": VERSION,
        "protocol": protocol(),
        "created_utc": now(),
        "first_retrieval_utc": first,
        "status": "NEW_CONFIRMATION_NOT_AVAILABLE",
        "new_confirmed_coffee_groups": 0,
        "C00_C01_confirmation_delta": None,
        "confirmation_model_runs": 0,
        "training_runs": 0,
        "zenodo_version_audit": versions,
        "zenodo_version_evidence": version_evidence,
        "new_targeted_routes": routes,
        "route_count": len(routes),
        "additional_actual_workbook_inspections": additional,
        "typed_actual_parses": [public_typed_summary(p) for p in parsed],
        "typed_private_artifacts": parsed_files,
        "prior_source_deduplication": dedup,
        "source_schema_scope": "Two typed sensory sources plus two additional original workbooks parsed substantively; typed availability is not a successful fixed-C00/C01 confirmation cohort. No unrelated attribute model fitted.",
        "next_required_external_input": "Rights-cleared original multicoffee observations with disjoint A/B/T and traceable unused coffee identities; owner freezes them before fixed candidate evaluation.",
    }
    # Reparse original sources while preserving the first actual receipt timestamp.
    report_path = owner / R7 / "confirmation_acquisition_summary.private.json"
    if report_path.exists():
        previous = json.loads(report_path.read_text())
        report["created_utc"] = previous["created_utc"]
        if report != previous:
            raise ValueError("SOURCE_REPARSE_DIFFERS_FROM_SEALED_REPORT")
    receipt = save(report_path, report)
    return {
        "status": report["status"],
        "summary_owner_relative_path": str(R7 / receipt["filename"]),
        "summary_sha256": receipt["sha256"],
        "routes": len(routes),
        "confirmation_groups": 0,
        "model_runs": 0,
    }


def verify(owner):
    owner = Path(owner)
    path = owner / R7 / "confirmation_acquisition_summary.private.json"
    report = json.loads(path.read_text())
    if report["protocol"] != protocol():
        raise ValueError("SOURCE_AUDIT_PROTOCOL_CHANGED")
    for item in report["typed_private_artifacts"]:
        checked_bytes(owner / item["owner_relative_path"], item["sha256"])
    evidence = report["zenodo_version_evidence"] + [
        x for route in report["new_targeted_routes"] for x in route["evidence"]
    ]
    evidence.append(report["prior_source_deduplication"]["remote_hemp_metadata"])
    for item in evidence:
        checked_bytes(owner / item["owner_relative_path"], item["sha256"])
    if (
        report["new_confirmed_coffee_groups"] != 0
        or report["confirmation_model_runs"]
        or report["training_runs"]
    ):
        raise ValueError("UNREGISTERED_CONFIRMATION_OR_TRAINING")
    return {
        "operation": "CACHED_SOURCE_HASH_VERIFICATION_NOT_REACQUISITION_OR_REFIT",
        "status": "PASS",
        "files_verified": len(evidence) + len(report["typed_private_artifacts"]),
        "summary_sha256": digest(path.read_bytes()),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--owner", required=True, type=Path)
    parser.add_argument("--phase", choices=("parse", "verify"), default="verify")
    args = parser.parse_args()
    print(
        json.dumps(
            build_report(args.owner) if args.phase == "parse" else verify(args.owner),
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
