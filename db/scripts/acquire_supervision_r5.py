"""Persistent, typed R5 source acquisition. No model fitting or label conflation."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import datetime
import hashlib
import io
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import tarfile
import tempfile
import urllib.error
import urllib.request
import urllib.parse
import zipfile

CATA = [
    "Tea.floral",
    "Fruit",
    "Citrus",
    "Green.veg",
    "Paper.wood",
    "Burnt",
    "Cereal",
    "Nutty",
    "Dark.chocolate",
    "Caramel",
    "Bitter",
    "Astringent",
    "Roasted",
    "Sour",
    "Thick.viscous",
    "Sweet",
    "Rubber",
]
JAR = ["Temp", "Flavor.intensity", "Acidity", "Mouthfeel"]
COTTER_SHA = "931aff6185381d5079bf93c4727bbbe65ff58ecfb524d2d3b6046eead2009114"
VERSION = "m2-r5-typed-source-acquisition.v1"
FIRSTBLOOM_REVISION = "a6cb0026d1af9642724793c799bbc48dc189ba35"


def now():
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


def digest(data):
    return hashlib.sha256(data).hexdigest()


def safe_source_url(url):
    """Keep source identity, never persist temporary redirect credentials."""
    parsed = urllib.parse.urlsplit(url)
    sensitive = {"token", "signature", "key-pair-id", "expires"}
    query = [
        (k, v)
        for k, v in urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
        if k.lower() not in sensitive and not k.lower().startswith("x-amz-")
    ]
    return urllib.parse.urlunsplit(
        (parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(query), "")
    )


def save(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    raw = (
        value
        if isinstance(value, bytes)
        else (
            json.dumps(
                value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False
            )
            + "\n"
        ).encode()
    )
    if path.exists() and path.read_bytes() != raw:
        raise ValueError("PRESERVE_EXISTING_SOURCE_ARTIFACT:" + path.name)
    if not path.exists():
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, "wb") as stream:
            stream.write(raw)
    return {"filename": path.name, "sha256": digest(raw), "bytes": len(raw)}


def fetch(url, directory, filename, timeout=45):
    """One public request; no access-control bypass or unbounded retry loop."""
    path = Path(directory) / filename
    receipt = path.with_name(path.name + ".acquisition.json")
    if path.exists() and receipt.exists():
        previous = json.loads(receipt.read_text())
        if digest(path.read_bytes()) != previous["sha256"]:
            raise ValueError("SOURCE_CACHE_HASH_CHANGED")
        return previous
    request = urllib.request.Request(
        url, headers={"User-Agent": "CoffeeFlavorResearch/1.0 (academic source audit)"}
    )
    start = now()
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            raw = response.read()
            resolved = urllib.parse.urlsplit(response.url)
            # Public storage redirects may carry short-lived access signatures.
            # They are unnecessary provenance and never belong in audit output.
            canonical_resolved = urllib.parse.urlunsplit(
                (resolved.scheme, resolved.netloc, resolved.path, "", "")
            )
            result = {
                "requested_url": safe_source_url(url),
                "resolved_url": canonical_resolved,
                "status": response.status,
                "content_type": response.headers.get("Content-Type"),
                "retrieved_utc": start,
                **save(path, raw),
            }
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError) as error:
        result = {
            "requested_url": safe_source_url(url),
            "retrieved_utc": start,
            "status": getattr(error, "code", "NETWORK_ERROR"),
            "error_type": type(error).__name__,
        }
    save(receipt, result)
    return result


def zip_inventory(path):
    """Inspect nested public supplements without treating an empty form as data."""
    rows = []

    def inspect(raw, prefix, depth=0):
        if depth > 3:
            raise ValueError("NESTED_ARCHIVE_BOUND_EXCEEDED")
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                if (
                    member.file_size > 100_000_000
                    or sum(m.file_size for m in archive.infolist()) > 300_000_000
                ):
                    raise ValueError("ARCHIVE_INSPECTION_SIZE_BOUND")
                content = archive.read(member)
                identity = prefix + member.filename
                rows.append(
                    {
                        "member": identity,
                        "bytes": len(content),
                        "sha256": digest(content),
                    }
                )
                if member.filename.lower().endswith(".zip"):
                    inspect(content, identity + "!", depth + 1)

    inspect(Path(path).read_bytes(), "")
    return rows


def parse_cotter(path):
    raw = Path(path).read_bytes()
    records = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    if not records or any(
        name not in records[0]
        for name in ["Judge", "Brew", *CATA, *JAR, "Liking", "Purchase.intent"]
    ):
        raise ValueError("COTTER_NATIVE_SCHEMA_REQUIRED")
    identities = [(row["Judge"], row["Brew"]) for row in records]
    if len(identities) != len(set(identities)):
        raise ValueError("DUPLICATE_PARTICIPANT_CONDITION_OBSERVATION")
    parsed = []
    for index, row in enumerate(records, 2):
        binary = {}
        for name in CATA:
            if row[name] not in {"0", "1"}:
                raise ValueError("CATA_MISSING_OR_NONBINARY_NOT_A_ZERO")
            binary[name] = {
                "value": int(row[name]),
                "protocol": "COMPLETE_SHOWN_CATA_BALLOT",
                "status": (
                    "SELECTED"
                    if row[name] == "1"
                    else "NOT_SELECTED_WITHIN_THIS_BALLOT"
                ),
                "absolute_sensory_absence": False,
                "is_intensity": False,
            }
        jar = {}
        for name in JAR:
            value = int(row[name])
            if value not in range(1, 6):
                raise ValueError("NATIVE_JAR_SCALE_OUTSIDE_1_TO_5")
            jar[name] = {
                "value": value,
                "protocol": "ADEQUACY_JAR_1_TO_5",
                "middle": 3,
                "is_perceived_intensity": False,
                "is_quality": False,
            }
        liking = int(row["Liking"])
        purchase = int(row["Purchase.intent"])
        if liking not in range(1, 10) or purchase not in range(1, 6):
            raise ValueError("NATIVE_HEDONIC_OR_PURCHASE_SCALE_MISMATCH")
        t, tds, extraction = map(float, row["Brew"].split("-"))
        parsed.append(
            {
                "record_id": "cotter:original-row:" + str(index),
                "collection_doi": "10.25338/B8993H",
                "participant_id": "cotter:judge:" + row["Judge"],
                "condition_id": "cotter:brew:" + row["Brew"],
                "coffee_material_id": "cotter:honduras_single_medium_roast_material",
                "session_id": "cotter:session:" + row["Session Number"],
                "production_context": {
                    "target_temperature_c": t,
                    "target_tds_percent": tds,
                    "target_extraction_percent": extraction,
                    "preparation_family": "filter_percolation",
                    "source_native_roast": "medium",
                    "mapped_seven_level_roast": None,
                },
                "cata": binary,
                "jar": jar,
                "liking": {"value": liking, "protocol": "HEDONIC_1_TO_9"},
                "purchase_intent": {"value": purchase, "protocol": "BIPOLAR_1_TO_5"},
                "quality": None,
                "measured_attribute_intensity": None,
                "task_roles": {
                    "source_native_binary_cata": True,
                    "same_participant_paired_views": True,
                    "cross_coffee_material_generalization": False,
                    "fine_descriptor_gold": False,
                    "JAR_as_intensity": False,
                    "liking_as_sensory_target": False,
                },
                "raw_source_row": index,
            }
        )
    judges, brews = {r["Judge"] for r in records}, {r["Brew"] for r in records}
    complete = len(records) == len(judges) * len(brews)
    return parsed, {
        "version": VERSION,
        "source_sha256": digest(raw),
        "same_as_r1_full_csv": digest(raw) == COTTER_SHA,
        "actual_observations": len(records),
        "participants": len(judges),
        "production_conditions": len(brews),
        "participant_condition_complete_cartesian_matrix": complete,
        "coffee_material_groups": 1,
        "sessions": len({r["Session Number"] for r in records}),
        "session_condition_pairs": len(
            {(r["Session Number"], r["Brew"]) for r in records}
        ),
        "native_columns": len(records[0]),
        "complete_cata_columns": len(CATA),
        "complete_cata_cells": len(records) * len(CATA),
        "jar_columns": JAR,
        "source_native_intensity_columns": 0,
        "liking_columns": 1,
        "quality_columns": 0,
        "cata_source_zero_cells": sum(r[c] == "0" for r in records for c in CATA),
        "metadata_3168_vs_data_3186": (
            "SOURCE_ABSTRACT_TYPO_3168; ACTUAL_CSV_AND_118_TIMES_27_EQUAL_3186"
            if len(records) == 3186
            else "VERSION_REQUIRES_REVIEW"
        ),
        "new_collection_sources_over_prior_rounds": (
            0 if digest(raw) == COTTER_SHA else None
        ),
        "new_coffee_material_groups_over_prior_rounds": (
            0 if digest(raw) == COTTER_SHA else None
        ),
        "permitted_role": "SOURCE_NATIVE_PAIRED_CONSUMER_VIEW; SAME_ONE_MATERIAL_WITH_PARTICIPANT_AND_CONDITION_ISOLATION",
        "not_permitted_role": "27_INDEPENDENT_BEANS_OR_JAR_TRUE_INTENSITY_OR_ABSOLUTE_ABSENCE",
    }


def read_csv(path):
    return list(
        csv.DictReader(io.StringIO(Path(path).read_bytes().decode("utf-8-sig")))
    )


def parse_firstbloom(directory):
    """Preserve actual review observations; never attach aggregate tags to people.

    Review IDs identify observation rows, not people. Product IDs group releases;
    raw-lot and participant overlap remain unknown. No A/B/T is fabricated here.
    """
    directory = Path(directory)
    reviews = read_csv(
        directory / "firstbloom_first_bloom_user_product_reviews_202312201402.csv"
    )
    products = read_csv(directory / "firstbloom_product_releases.csv")
    release_map = {row["product_release_id"]: row for row in products}
    if len(release_map) != len(products):
        raise ValueError("NONUNIQUE_FIRSTBLOOM_RELEASE_JOIN")
    if len({r["id"] for r in reviews}) != len(reviews):
        raise ValueError("NONUNIQUE_FIRSTBLOOM_REVIEW_OBSERVATION")
    required = {
        "id",
        "review_text",
        "rating",
        "product_release_id",
        "body_intensity",
        "acidity_intensity",
        "sweetness_intensity",
        "finish_intensity",
    }
    if not reviews or not required.issubset(reviews[0]):
        raise ValueError("FIRSTBLOOM_NATIVE_REVIEW_SCHEMA_REQUIRED")
    parsed = []
    by_release, by_product = defaultdict(list), defaultdict(list)
    score_zeros = Counter()
    for row in reviews:
        if row["product_release_id"] not in release_map:
            raise ValueError("FIRSTBLOOM_REVIEW_WITHOUT_ACTUAL_PRODUCT_JOIN")
        product = release_map[row["product_release_id"]]
        native = {}
        for field in sorted(
            required
            & {
                "body_intensity",
                "acidity_intensity",
                "sweetness_intensity",
                "finish_intensity",
            }
        ):
            value = None if not row[field].strip() else float(row[field])
            if value is not None and (not math.isfinite(value) or not 0 <= value <= 5):
                raise ValueError("FIRSTBLOOM_NATIVE_SCORE_OUTSIDE_OBSERVED_0_TO_5")
            score_zeros[field] += value == 0
            native[field] = {
                "raw_value": value,
                "value": value if value not in {None, 0} else None,
                "status": (
                    "ZERO_OR_DEFAULT_UNRESOLVED_MASKED"
                    if value == 0
                    else "MISSING" if value is None else "OBSERVED_NATIVE_SCORE"
                ),
                "anchors_verified": False,
                "absolute_sensory_absence": False,
                "eligible_for_true_intensity_loss": False,
            }
        text = row["review_text"].strip()
        parsed.append(
            {
                "observation_id": "firstbloom:review:" + row["id"],
                "source_group_id": "firstbloom:product:" + product["product_id"],
                "release_id": "firstbloom:release:" + row["product_release_id"],
                "participant_id": None,
                "raw_lot_id": None,
                "collection": "firstbloom-data",
                "collection_version": FIRSTBLOOM_REVISION,
                "review_text": text or None,
                "review_text_sha256": digest(text.encode()) if text else None,
                "rating": {
                    "value": float(row["rating"]),
                    "task": "SOURCE_STAR_EVALUATION; HEDONIC_OR_QUALITY_NOT_DISAMBIGUATED",
                    "sensory_target": False,
                },
                "native_score_observations": native,
                "cata": None,
                "jar": None,
                "actual_question_option_exposure": None,
                "source_C0": None,
                "source_C1": None,
                "source_created_at": row.get("created_at"),
                "task_roles": {
                    "positive_text_nonmention_is_unknown": True,
                    "separate_source_observation_unit": True,
                    "independent_participant_verified": False,
                    "blind_sensory_panel": False,
                    "professional_gold": False,
                    "runtime_answer_trajectory": False,
                },
            }
        )
        if text:
            by_release[row["product_release_id"]].append(text)
            by_product[product["product_id"]].append(text)
    aggregate_tags = read_csv(
        directory / "firstbloom_product_release_tasting_notes.csv"
    )
    # These sets belong to a release-level aggregation, not a known review row.
    aggregate_tag_views = [
        {
            "release_id": "firstbloom:release:" + row["product_release_id"],
            "source_group_id": "firstbloom:product:"
            + release_map[row["product_release_id"]]["product_id"],
            "roaster_tag_ids_native": row["roaster_tasting_notes"],
            "user_tag_ids_native": row["user_review_tasting_notes"],
            "unit": "RELEASE_AGGREGATE_WITH_UNKNOWN_CONTRIBUTOR_MEMBERSHIP",
            "individual_cata_ballot": False,
            "independent_of_review_text": None,
        }
        for row in aggregate_tags
        if row["product_release_id"] in release_map
    ]
    all_text = [r["review_text"] for r in parsed if r["review_text"]]
    summary = {
        "version": VERSION,
        "collection_version": FIRSTBLOOM_REVISION,
        "license": "CC-BY-4.0",
        "attribution": "Firstbloom Data by Alex Caza; private typed extraction and source-group audit",
        "observations": len(parsed),
        "review_observation_ids_unique": True,
        "distinct_participants": None,
        "participant_ids_present": False,
        "reviewed_release_groups": len({r["release_id"] for r in parsed}),
        "reviewed_product_groups": len({r["source_group_id"] for r in parsed}),
        "verified_independent_raw_lots": None,
        "nonempty_text_observations": len(all_text),
        "exact_unique_nonempty_texts": len(set(all_text)),
        "nonempty_text_release_groups": len(by_release),
        "nonempty_text_product_groups": len(by_product),
        "release_groups_with_at_least_3_distinct_text_units": sum(
            len(set(v)) >= 3 for v in by_release.values()
        ),
        "product_groups_with_at_least_3_distinct_text_units": sum(
            len(set(v)) >= 3 for v in by_product.values()
        ),
        "release_nonempty_text_count_histogram": dict(
            sorted(Counter(map(len, by_release.values())).items())
        ),
        "unresolved_zero_cells_masked": dict(score_zeros),
        "true_intensity_target_cells_admitted": 0,
        "aggregate_user_tag_release_groups": sum(
            r["user_tag_ids_native"] not in {"", "{}"} for r in aggregate_tag_views
        ),
        "aggregate_tags_copied_to_individual_reviews": 0,
        "new_collection_sources": 0,
        "new_unused_observation_view": True,
        "overlap_with_prior_model_coffee_groups": "REQUIRES_EXISTING_SOURCE_ID_AUDIT; NO_NEW_COFFEE_COUNT_CLAIM",
        "role": "WEAK_REVIEW_OBSERVATION_PILOT; NO_BLIND_OR_PARTICIPANT_INDEPENDENCE_CLAIM",
        "limitation": "Self-selected reviews, possible shared authors and marketing exposure, missing brew/roast, unknown lot overlap; native scale anchors and zero defaults unresolved.",
    }
    return parsed, aggregate_tag_views, summary


def workbook_inventory(path):
    """Inspect actual row/cell structure without inferring participants from rows."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return [
        {
            "sheet": sheet.title,
            "rows_including_headers_and_blank": sheet.max_row,
            "columns": sheet.max_column,
            "nonempty_rows": sum(
                any(v is not None for v in row) for row in sheet.values
            ),
        }
        for sheet in workbook
    ]


def parse_peru_quality(path):
    """Native cupping quality repeat rows; attribute names do not imply intensity."""
    import openpyxl

    sheet = openpyxl.load_workbook(path, data_only=True, read_only=True)["Catacion"]
    rows = list(sheet.values)
    headers = [str(v).strip() if v is not None else "" for v in rows[1]]
    expected = [
        "Time",
        "Treatments",
        "Repetition",
        "Aroma",
        "Taste",
        "Acidity",
        "Body",
        "Balance",
        "Overall",
        "Sweetness",
        "Overtaste",
        "Uniformity",
        "clean cup",
    ]
    if headers[1:] != expected:
        raise ValueError("PERU_NATIVE_CUPPING_SCHEMA_REQUIRED")
    parsed, identities = [], set()
    for row_number, raw in enumerate(rows[2:], 3):
        if all(v is None for v in raw):
            continue
        row = dict(zip(headers, raw))
        identity = (row["Time"], row["Treatments"], row["Repetition"])
        if identity in identities:
            raise ValueError("DUPLICATE_PERU_NATIVE_REPEAT")
        identities.add(identity)
        quality = {}
        for field in expected[3:]:
            value = row[field]
            if value is not None and (
                not isinstance(value, (int, float))
                or not math.isfinite(value)
                or not 0 <= value <= 10
            ):
                raise ValueError("PERU_NATIVE_QUALITY_SCALE_MISMATCH")
            quality[field] = {
                "value": value,
                "task": "CUPPING_QUALITY",
                "is_intensity": False,
                "unobserved_is_zero": False,
            }
        parsed.append(
            {
                "observation_id": "peru:cupping-row:" + str(row_number),
                "collection_doi": "10.6084/m9.figshare.31895233.v2",
                "coffee_material_group": "peru:la_peca_one_farm_june2021_harvest",
                "production_condition_id": "peru:"
                + row["Time"]
                + ":"
                + row["Treatments"],
                "native_repeat_id": row["Repetition"],
                "participant_id": None,
                "repeat_is_participant_verified": False,
                "quality": quality,
                "cata": None,
                "jar": None,
                "intensity": None,
                "sensory_descriptors": None,
                "task_role": "QUALITY_ONLY; NO_FINE_DESCRIPTOR_OR_INTENSITY_SUPERVISION",
            }
        )
    counts = Counter(r["production_condition_id"] for r in parsed)
    return parsed, {
        "source_doi": "10.6084/m9.figshare.31895233.v2",
        "source_sha256": digest(Path(path).read_bytes()),
        "license": "CC-BY-4.0",
        "actual_quality_rows": len(parsed),
        "production_conditions": len(counts),
        "condition_row_count_histogram": dict(sorted(Counter(counts.values()).items())),
        "coffee_material_groups": 1,
        "native_quality_dimensions": 10,
        "sensory_intensity_cells": 0,
        "participant_ids_verified": False,
        "actual_rows_not_paper_multiplication": True,
        "count_discrepancy": "FILE26; PAPER_TEXT28; NOMINAL5_TIMES2_TIMES3_EQUALS30; DO_NOT_IMPUTE_MISSING_CONTROL_REPLICATES",
        "role": "QUALITY_ONLY_NOT_ELIGIBLE_FOR_CURRENT_SENSORY_TARGET",
    }


def parse_blendstat(directory):
    """Source-native aggregate Flavour/Bitterness, with conservative design groups.

    Co-holding equal blend vectors across processing/concentration is a split
    precaution, not a claim that the observations share an actual coffee lot.
    """
    directory = Path(directory)
    records, file_hashes = [], {}
    for process, name in [("PEELED_CHERRY", "DataCD"), ("NATURAL", "DataNAT")]:
        path = directory / ("blendstat_" + name + ".csv")
        rows = read_csv(path)
        required = {
            "Sample",
            "Exp",
            "CEB",
            "CT",
            "CC",
            "CEA",
            "Conc",
            "Flavour",
            "Bitterness",
            "Score",
            "Body",
            "Acidity",
        }
        if not rows or not required.issubset(rows[0]):
            raise ValueError("BLENDSTAT_NATIVE_SCHEMA_REQUIRED")
        identities = [(r["Exp"], r["Sample"]) for r in rows]
        if len(identities) != len(set(identities)):
            raise ValueError("DUPLICATE_BLENDED_CONDITION_IDENTITY")
        file_hashes[path.name] = digest(path.read_bytes())
        for row in rows:
            mix = [float(row[k]) for k in ["CEB", "CT", "CC", "CEA"]]
            if (
                any(not math.isfinite(v) or not 0 <= v <= 1 for v in mix)
                or abs(sum(mix) - 1) > 0.011
            ):
                raise ValueError("INVALID_OBSERVED_BLEND_VECTOR")
            concentration = float(row["Conc"])
            if concentration not in {0.07, 0.1}:
                raise ValueError("UNSUPPORTED_BLEND_CONCENTRATION")
            native = {}
            for key in ["Flavour", "Bitterness", "Body", "Acidity", "Score"]:
                value = float(row[key])
                if not math.isfinite(value) or not 0 <= value <= 10:
                    raise ValueError("BLENDSTAT_NATIVE_SCORE_RANGE")
                known_sensory = key in {"Flavour", "Bitterness"}
                native[key] = {
                    "value": value,
                    "source_native_scale": [0, 10],
                    "task": (
                        "SOURCE_NATIVE_SENSORY_RATING"
                        if known_sensory
                        else (
                            "OVERALL_QUALITY"
                            if key == "Score"
                            else "ATTRIBUTE_RATING_CONSTRUCT_UNRESOLVED"
                        )
                    ),
                    "eligible_minimal_flavour_to_bitterness_task": known_sensory,
                    "physiological_interval_scale_claim": False,
                    "zero_is_absolute_absence": False,
                }
            group = "blendstat:coheld-composition:" + "-".join(
                format(v, ".2f") for v in mix
            )
            records.append(
                {
                    "observation_id": "blendstat:"
                    + name
                    + ":"
                    + row["Exp"]
                    + ":"
                    + row["Sample"],
                    "collection": "BLENDSTAT_MANTIQUEIRA_PROJECT_304974_2015_3",
                    "package_version": "1.0.6",
                    "source_file": path.name,
                    "process": process,
                    "native_experiment": row["Exp"],
                    "native_sample": row["Sample"],
                    "native_mix": dict(zip(["CEB", "CT", "CC", "CEA"], mix)),
                    "native_concentration_wv": concentration,
                    "conservative_split_group": group,
                    "split_group_is_independent_coffee_material": False,
                    "participant_id": None,
                    "raw_lot_id": None,
                    "observation_unit": "SOURCE_CONDITION_AGGREGATE; NO_INDIVIDUAL_JUDGE_ROWS",
                    "native_ratings": native,
                    "minimum_auxiliary_view": {
                        "A": [],
                        "B": {"native.Flavour": native["Flavour"]["value"]},
                        "T": {"native.Bitterness": native["Bitterness"]["value"]},
                    },
                    "available_runtime_question_answers": False,
                    "B_T_are_separate_recorded_fields_not_independent_judges": True,
                    "forbidden_predictors": [
                        "Score",
                        "Body",
                        "Acidity",
                        "native_mix",
                        "native_concentration_wv",
                        "process",
                        "native_experiment",
                        "native_sample",
                        "conservative_split_group",
                    ],
                    "task_role": "STATIC_SOURCE_VALUE_DIAGNOSTIC; NOT_PRODUCT_Q3_OR_DESCRIPTOR_TARGET",
                }
            )
    summary = {
        "version": VERSION,
        "package": "Blendstat",
        "package_version": "1.0.6",
        "license": "GPL-3",
        "data_license_evidence": "AUTHOR_RELEASED_CRAN_PACKAGE_DESCRIPTION; EXPLICIT_DATA_RD; NO_SEPARATE_DATA_EXCEPTION",
        "article_doi": "10.5935/1806-6690.20190041",
        "source_file_sha256": file_hashes,
        "actual_aggregate_rows": len(records),
        "processing_views": 2,
        "rows_by_processing": dict(Counter(r["process"] for r in records)),
        "unique_process_mix_concentration_conditions": len(
            {
                (
                    r["process"],
                    r["conservative_split_group"],
                    r["native_concentration_wv"],
                )
                for r in records
            }
        ),
        "conservative_composition_groups": len(
            {r["conservative_split_group"] for r in records}
        ),
        "measured_participant_rows": 0,
        "participant_count_in_paper_not_rows": 5,
        "independent_coffee_material_count": None,
        "actual_CSV_experiment_map": {
            "1": "BOURBON_0.07",
            "2": "ACAIA_0.07",
            "3": "BOURBON_0.10",
            "4": "ACAIA_0.10",
        },
        "paper_table1_experiment2_3_disagree_with_actual_CSV": True,
        "actual_input": "native.Flavour",
        "actual_target": "native.Bitterness",
        "source_native_scale": [0, 10],
        "input_target_same_aggregate_observation": True,
        "excluded_quality": ["Score"],
        "excluded_unclear_attribute_constructs": ["Body", "Acidity"],
        "new_collection_sources": 1,
        "new_independent_coffee_materials_claim": None,
        "role": "ADMIT_MINIMAL_STATIC_SOURCE_VALUE_DIAGNOSTIC_ONLY",
        "limitation": "Same small constituent pool; 72 aggregate rows, not 72 beans or judges; composition cohold is a conservative design split, not a source identity join. No fixed physiological interval or independent-view claim.",
    }
    return records, summary


def extract_blendstat(directory, rscript=None):
    """Read packaged numeric data only; never install or execute package code."""
    directory = Path(directory)
    executable = rscript or shutil.which("Rscript")
    if not executable:
        raise RuntimeError("R_SCRIPT_REQUIRED_TO_READ_ORIGINAL_RDA_OBJECTS")
    allowed = {
        "Blendstat/data/DataCD.rda",
        "Blendstat/data/DataNAT.rda",
        "Blendstat/man/DataCD.Rd",
        "Blendstat/man/DataNAT.Rd",
        "Blendstat/DESCRIPTION",
    }
    inventory = []
    with tarfile.open(directory / "Blendstat_1.0.6.tar.gz") as archive:
        for member in archive.getmembers():
            if member.name in allowed:
                if not member.isfile() or member.size > 1_000_000:
                    raise ValueError("BLENDSTAT_DATA_ARCHIVE_MEMBER_INVALID")
                inventory.append(
                    save(
                        directory / ("blendstat_" + Path(member.name).name),
                        archive.extractfile(member).read(),
                    )
                )
    if len(inventory) != len(allowed):
        raise ValueError("BLENDSTAT_EXPECTED_NATIVE_MEMBERS_NOT_PRESENT")
    r_code = r"""
args <- commandArgs(trailingOnly=TRUE)
for (n in c('DataCD','DataNAT')) {
  e <- new.env(); load(file.path(args[1],paste0('blendstat_',n,'.rda')),envir=e)
  d <- get(n,envir=e)
  stopifnot(is.data.frame(d), nrow(d) == 36L, ncol(d) == 12L,
            all(vapply(d,is.numeric,logical(1))),
            identical(names(d), c('Sample','Exp','CEB','CT','CC','CEA',
                                  'Conc','Body','Flavour','Acidity','Bitterness','Score')))
  write.csv(d,file.path(args[2],paste0('blendstat_',n,'.csv')),row.names=FALSE)
}
"""
    with tempfile.TemporaryDirectory() as scratch:
        result = subprocess.run(
            [str(executable), "-", str(directory), scratch],
            input=r_code,
            text=True,
            capture_output=True,
            timeout=30,
        )
        if result.returncode:
            raise ValueError("RDA_NUMERIC_EXTRACTION_FAILED:" + result.stderr[-500:])
        for name in ["blendstat_DataCD.csv", "blendstat_DataNAT.csv"]:
            inventory.append(
                save(directory / name, (Path(scratch) / name).read_bytes())
            )
    return {
        "extracted_artifacts": inventory,
        "package_code_executed": False,
        "model_fits": 0,
    }


def parse_great_american(path):
    """Minimize the original anonymous survey; rights unresolved means no fit."""
    rows = read_csv(path)
    if not rows or "Submission ID" not in rows[0]:
        raise ValueError("GACTT_NATIVE_SUBMISSION_REQUIRED")
    if len({r["Submission ID"] for r in rows}) != len(rows):
        raise ValueError("DUPLICATE_GACTT_SUBMISSION")
    parsed = []
    for row in rows:
        for coffee in "ABCD":
            ratings = {}
            for field in ["Bitterness", "Acidity", "Personal Preference"]:
                raw = row[f"Coffee {coffee} - {field}"].strip()
                value = int(raw) if raw else None
                if value is not None and value not in range(1, 6):
                    raise ValueError("GACTT_OBSERVED_RATING_OUTSIDE_1_TO_5")
                ratings[field] = {
                    "value": value,
                    "status": "MISSING" if value is None else "OBSERVED",
                    "construct": (
                        "PERSONAL_PREFERENCE"
                        if field == "Personal Preference"
                        else "NATIVE_SENSORY_RATING; ANCHORS_PENDING"
                    ),
                    "missing_is_zero": False,
                }
            notes = row[f"Coffee {coffee} - Notes"].strip() or None
            parsed.append(
                {
                    "observation_id": "gactt:" + row["Submission ID"] + ":" + coffee,
                    "anonymous_submission_id": "gactt:" + row["Submission ID"],
                    "source_sample_id": "gactt:" + coffee,
                    "raw_coffee_material_id": None,
                    "person_identity_independence_verified": False,
                    "native_ratings": ratings,
                    "native_notes": notes,
                    "empty_native_notes_is_sensory_absence": False,
                    "cata_ballot": None,
                    "jar": None,
                    "source_C0": None,
                    "source_C1": None,
                    "rights_status": "ORIGINAL_AUTHOR_PUBLIC_INSPECTION_LINK; TRAINING_AND_REDISTRIBUTION_PERMISSION_NOT_YET_EXPLICIT",
                    "training_allowed": False,
                }
            )
    return parsed, {
        "source_sha256": digest(Path(path).read_bytes()),
        "actual_unique_submissions": len(rows),
        "actual_source_columns": len(rows[0]),
        "four_sample_observation_slots": len(parsed),
        "observations_with_both_sensory_ratings": sum(
            all(
                r["native_ratings"][k]["value"] is not None
                for k in ["Bitterness", "Acidity"]
            )
            for r in parsed
        ),
        "observations_with_notes": sum(bool(r["native_notes"]) for r in parsed),
        "presented_sample_codes": 4,
        "independent_raw_coffee_materials": None,
        "training_allowed": False,
        "new_admitted_sources": 0,
        "rights_status": "HOLD_ORIGINAL_REUSE_TERMS; THIRD_PARTY_KAGGLE_LICENSE_NOT_RELIED_ON",
        "unrelated_demographic_fields_in_typed_extraction": 0,
        "role": "ACTUAL_PAIRED_SURVEY_ACQUIRED_BUT_NOT_ADMITTED_FOR_TRAINING",
    }


def source_ledger():
    """Compact reviewed route outcomes, not a count of successful new datasets."""
    columns = [
        "paper_id",
        "supported_methodological_point",
        "data_unit",
        "applicability",
        "limitation",
        "resulting_experiment_change",
    ]
    entries = [
        (
            "10.25338/B8993H;10.1038/s41538-026-00779-7",
            "CATA selection and JAR adequacy differ from intensity and liking",
            "3186 judge-condition rows;118 judges;27 conditions;1 material",
            "Existing CC0 source; current v4 CSV exactly matches R1 FULL hash",
            "Metadata also says3168; 118x27 and actualCSV3186; later AI paper same collection",
            "Reuse native CATA only;0 new sources/beans;unselected is ballot response,not absolute absence",
        ),
        (
            "10.3390/foods14040593",
            "Repeated professional evaluation requires actual individual data",
            "Acquired XML plus1-page roast-settings TableS1 PDF",
            "Method evidence only; same R2 source/member hash",
            "Individual espresso responses restricted for privacy/ethics; paper people not acquiredrows",
            "No new training rows; do not invent repeated paired views",
        ),
        (
            "10.3390/foods15040678",
            "A scorecard schema is not a filled observation",
            "Acquired XML plus1-page PDF of6 blank scorecards",
            "Method evidence only; same R2 inner PDF hash",
            "Outer ZIP changed; underlying form unchanged; JAR/liking fields remain distinct",
            "0 filled observations; no scorecard-derived targets",
        ),
        (
            "S0963996925000067",
            "Public paper availability does not authorize raw-data reuse",
            "No new request or file acquisition",
            "STOP_USER_REPORTED_NO_SHARING_RIGHTS",
            "Rights unavailable per user; no retry, bypass, contact, or payment",
            "Exclude; retain stop outcome only",
        ),
        (
            "10.1371/journal.pone.0155845;10.17026/DANS-ZKE-2WGQ",
            "Expert language can differ by modality and expertise without being unique truth",
            "Existing Croijmans coffee smell/flavor responses",
            "Reuse prior licensed collection and identity-linked views",
            "Already R2/R3; not new people or coffee; comments are reported perceptions",
            "Keep native smell/flavor pairing and expert-consumer distinction;0 new collection",
        ),
        (
            "10.17632/3yv7bdrczd.1",
            "Information treatment can alter consumer evaluations",
            "Actual xlsx93 rows;90 unique IDs;3 IDs repeated twice;1 coffee",
            "CC BY4 actual file acquired and inspected",
            "6 rows have ambiguous repeatedID; native scale anchors unresolved; treatment is not coffee identity",
            "Quarantine sensory-rating fit; no93 independent-person or3-coffee claim",
        ),
        (
            "firstbloom-data@a6cb0026d1af9642724793c799bbc48dc189ba35",
            "Reviews are observations; release-level tag sets are not individual ballots",
            "1231 reviews;704 nonblank;805 releases/768 products;no participantID",
            "Founder explicitly licenses data CC BY4; unused view of old collection",
            "Marketing source rights separate; score zero/anchors unresolved; author and lot overlap unknown",
            "Typed weak observation audit only; no new collection; aggregate tags never assigned to individuals",
        ),
        (
            "10.6084/m9.figshare.31895233.v2;10.1093/ijfood/vvag095",
            "Cupping quality and literature odor annotations are different target types",
            "26 actual repeat-quality rows;10 condition-level TableS8 descriptions;1 farm harvest",
            "CC BY4 raw xlsx and official DOCX acquired",
            "Paper says28 and nominal design30; real controls single rows; R1-R3 biological repeats,not graders",
            "Preserve26; no imputation; chemical TableS6 odor labels excluded; no independent view claim",
        ),
        (
            "10.6084/m9.figshare.28333775.v1",
            "Marginal question summaries cannot reconstruct respondent-level pairing",
            "Actual summary workbook90x12; response count tables for2 coffee types",
            "CC BY4 source workbook inspected",
            "No participant-response matrix; predictive-focus appreciation differs from sensory truth",
            "No fabricated cross-question pairs or main-task admission",
        ),
        (
            "10.6084/m9.figshare.28513958.v1;10.3390/su17052152",
            "Sensory quality attributes are not perceived intensities",
            "Workbook4 orange-peel treatment aggregate sensory profiles plus laboratory sheets",
            "CC BY4 workbook and primary XML acquired",
            "No individual graders; altered beverage conditions share coffee material",
            "Quality/laboratory observations separated; no current descriptor supervision",
        ),
        (
            "10.5281/zenodo.21961040",
            "Dataset title alone does not prove sensory data are present",
            "All5 actual files inspected;3CSV+2ODS with11 laboratory sheets",
            "CC BY4 original archive completely acquired",
            "Title mentions cupping; actual files contain pH/HPLC/GCMS, no human response table identified",
            "0 sensory rows admitted; chemical peak intensity never sensory intensity",
        ),
        (
            "Blendstat1.0.6;10.5935/1806-6690.20190041",
            "Native Flavour/Bitterness rating recovery can be tested separately from quality",
            "72 aggregate rows;32 processing-mixture identities;64 conditions;16 coheld mixture vectors",
            "Author-released GPL3 data; primary-paper scale0-10; shared author Cirillo",
            "No individual judge rows; small shared constituent pool; actual Exp2/3 differ from paper Table1",
            "Admit only actual Flavour to Bitterness static diagnostic;exclude Score/Body/Acidity/design features",
        ),
        (
            "10.1186/s12870-024-04890-3;10.18167/DVN1/NHCK5F",
            "Independent professional panels need their raw responses preserved",
            "4 genotype aggregate profiles; blank TableS1; supplementary chemistry",
            "CC BY4 primary XML/DOCX inspected; data API401/DataCite404",
            "8 judges in article; no acquired individual matrix; figure aggregates cannot be unpooled",
            "No raw-response fit; preserve protocol and unavailable-data outcome",
        ),
        (
            "GreatAmericanCoffeeTasteTest2023:bMOOQfeloH0",
            "Native notes/ratings can be paired by actual anonymous submission",
            "4042 unique submissions;16168 sample slots;15053 both sensory ratings;10005 notes",
            "Original author-linked CSV actually acquired and minimized",
            "Explicit training/redistribution terms and anchors not verified; raw material overlap unknown; third-party license insufficient",
            "Quarantine all fits; drop unrelated demographics;4 samplecodes not4042 coffees",
        ),
        (
            "10.3390/foods11030473",
            "Descriptive intensity and acceptance are separately measured",
            "Primary XML states12 concentrates and duplicate10-judge analyses; supplement contains figures only",
            "Article and actual supplement acquired",
            "Raw observations available on request pending privacy/ethics; no response matrix acquired",
            "Method evidence only; no paper-count expansion or synthetic pairs",
        ),
        (
            "10.1038/s41598-025-99921-w",
            "RATA1-7 includes intensity while quality score remains separate",
            "Current public article and exact-title Dryad query0 results",
            "Previously investigated source; primary protocol rechecked",
            "49 graders/67 samples in paper not public raw rows; private share link not used",
            "0 new observations; do not copy full wheel or create paper-derived labels",
        ),
        (
            "10.5061/dryad.v15dv423h",
            "A public endpoint response may still not expose a dataset",
            "Actual API200 with not-viewable message",
            "Prior source endpoint rechecked",
            "No downloadable native sensory matrix;30 brew conditions not30 beans",
            "No fit; retain source availability failure",
        ),
        (
            "10.1371/journal.pone.0223280",
            "Bitter stimulus subqualities can change retronasal perception",
            "Primary article7 stimulus-coffee aggregate profiles; no per-panelist data file",
            "CC BY4 method/context evidence",
            "Same coffee extract with added bitter substances; not distinct raw coffees",
            "Keep context point; no new independent coffee/paired-response claim",
        ),
        (
            "10.1038/s41538-026-00832-5",
            "Large citizen-science panel counts require released records",
            "Actual primary article; data deferred to larger study",
            "Method evidence only",
            "No actual participant dataset acquired despite large published counts",
            "No count-based expansion; no author contact sent",
        ),
        (
            "figshare7390667",
            "An open license is insufficient if native files cannot be retrieved",
            "3 metadata-listed Conilon sensory tables; actual file URLs404",
            "Metadata/license audited",
            "No actual xls content acquired; genotype table aggregates not individual judges",
            "No data admission; keep retrieval failure",
        ),
        (
            "figshare14318575",
            "Consumer expectations and tasting are different views",
            "Metadata-listed article figures and2 aggregate tables",
            "Candidate inspected at original repository",
            "No individual response file; publication participants not data rows",
            "No main-task pairing from marginal summaries",
        ),
        (
            "figshare8292704",
            "Repository supplements may omit the research data carried by an author package",
            "Published tables/plots; actual data found separately in Blendstat",
            "Same Mantiqueira project; deduplicate collection",
            "No per-consumer file; paper/package not two independent sources",
            "Count only one Blendstat collection; preserve actual package data",
        ),
    ]
    return [dict(zip(columns, row)) for row in entries]


def ledger_tsv():
    rows = source_ledger()
    out = io.StringIO()
    writer = csv.DictWriter(
        out, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n"
    )
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--owner-dir", type=Path, required=True)
    parser.add_argument(
        "--phase",
        choices=[
            "fetch",
            "parse-cotter",
            "parse-firstbloom",
            "parse-peru",
            "parse-blendstat",
            "extract-blendstat",
            "parse-gactt",
            "audit-ledger",
            "workbook-inventory",
            "zip-inventory",
        ],
        required=True,
    )
    parser.add_argument("--url")
    parser.add_argument("--filename")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--rscript", type=Path)
    args = parser.parse_args()
    directory = args.owner_dir / "revisions/r5/sources"
    if args.phase == "fetch":
        result = fetch(args.url, directory, args.filename)
    elif args.phase == "parse-cotter":
        records, result = parse_cotter(args.input)
        result["typed_private_artifact"] = save(
            directory / "cotter_typed_source_rows.private.json", records
        )
        save(directory / "cotter_public_data_role_summary.private.json", result)
    elif args.phase == "parse-firstbloom":
        records, tags, result = parse_firstbloom(directory)
        result["typed_private_artifact"] = save(
            directory / "firstbloom_native_observations.private.json", records
        )
        result["aggregate_tags_artifact"] = save(
            directory / "firstbloom_release_aggregate_tags.private.json", tags
        )
        save(directory / "firstbloom_data_role_summary.private.json", result)
    elif args.phase == "parse-peru":
        records, result = parse_peru_quality(args.input)
        result["typed_private_artifact"] = save(
            directory / "peru_native_observations.private.json", records
        )
        save(directory / "peru_data_role_summary.private.json", result)
    elif args.phase == "parse-blendstat":
        records, result = parse_blendstat(directory)
        result["typed_private_artifact"] = save(
            directory / "blendstat_native_observations.private.json", records
        )
        save(directory / "blendstat_data_role_summary.private.json", result)
    elif args.phase == "extract-blendstat":
        result = extract_blendstat(directory, args.rscript)
    elif args.phase == "parse-gactt":
        records, result = parse_great_american(args.input)
        result["typed_private_artifact"] = save(
            directory / "gactt_native_observations.private.json", records
        )
        save(directory / "gactt_data_role_summary.private.json", result)
    elif args.phase == "audit-ledger":
        result = save(directory / "source_ledger.private.tsv", ledger_tsv())
        result["rows"] = len(source_ledger())
    elif args.phase == "workbook-inventory":
        result = workbook_inventory(args.input)
    else:
        result = zip_inventory(args.input)
    print(json.dumps(result, sort_keys=True, allow_nan=False))


if __name__ == "__main__":
    main()
